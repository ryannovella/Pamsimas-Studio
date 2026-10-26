import os
import re
import uuid
import docx
from datetime import datetime, timedelta
from typing import Dict, Any, List, Optional
from app.backend.docx_service import parse_document, lookup_location, fuzzy_match_village
from app.backend.logger import log_error
from app.backend.report_generator import parse_date_range_str, normalize_personel_title


def parse_indo_date_to_iso(text: str) -> str:
    """Mengonversi teks tanggal Indonesia (e.g. '18 September 2026') ke ISO '2026-09-18'."""
    if not text:
        return ""
    # Format range seperti '05 - 06 Agustus 2026', ambil hari pertama
    m_range = re.search(r"(\d{1,2})\s*[-–—]\s*(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", text)
    if m_range:
        day = int(m_range.group(1))
        m_str = m_range.group(3).lower()
        year = int(m_range.group(4))
    else:
        m = re.search(r"(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", text)
        if not m:
            return ""
        day = int(m.group(1))
        m_str = m.group(2).lower()
        year = int(m.group(3))

    indo_month_map = {
        'januari': 1, 'jan': 1, 'februari': 2, 'feb': 2, 'maret': 3, 'mar': 3,
        'april': 4, 'apr': 4, 'mei': 5, 'juni': 6, 'jun': 6, 'juli': 7, 'jul': 7,
        'agustus': 8, 'agt': 8, 'agust': 8, 'september': 9, 'sep': 9, 'sept': 9,
        'oktober': 10, 'okt': 10, 'november': 11, 'nov': 11, 'desember': 12, 'des': 12
    }
    if m_str in indo_month_map:
        month = indo_month_map[m_str]
        return f"{year:04d}-{month:02d}-{day:02d}"
    return ""


def get_hplus1_iso(iso_date: str) -> str:
    """Menghitung H+1 dari tanggal ISO YYYY-MM-DD."""
    if not iso_date:
        return ""
    try:
        parts = [int(p) for p in iso_date.split("-")]
        d = datetime(parts[0], parts[1], parts[2]) + timedelta(days=1)
        return d.strftime("%Y-%m-%d")
    except Exception:
        return ""


def inspect_and_import_monev_docx(docx_path: str) -> Dict[str, Any]:
    """
    Menganalisis dokumen Monev Word eksisting (.docx), mengekstrak data administrasi,
    jadwal dinas, tabel pelaksana, daftar desa beserta narasi capaian, dan galeri foto,
    serta menghasilkan laporan audit kelengkapan dokumen.
    """
    if not os.path.exists(docx_path):
        return {"success": False, "error": f"File tidak ditemukan: {docx_path}"}

    doc = docx.Document(docx_path)
    filename = os.path.basename(docx_path)
    clean_fn = re.sub(r"^inspect_[a-f0-9]+_", "", filename, flags=re.I)

    from app.backend.docx_service import KNOWN_LOCATIONS

    # 1. Judul, Kabupaten, dan Tahun
    kabupaten = ""
    tahun = "2026"
    for p in doc.paragraphs[:15]:
        t = p.text.strip()
        m_title = re.search(r"Monitoring dan Evaluasi Kegiatan Pamsimas\s+(?:Kabupaten|Kota)?\s*([A-Za-z\s]+?)\s+TA\s*(\d{4})", t, re.I)
        if m_title:
            kabupaten = m_title.group(1).strip()
            tahun = m_title.group(2).strip()
            break

    # Jika kabupaten belum ketemu di judul, cari pola "Kabupaten X" di seluruh teks dokumen
    if not kabupaten:
        for p in doc.paragraphs:
            m_kab_doc = re.search(r"\b(?:Kabupaten|Kab\.?)\s+([A-Za-z]+)", p.text, re.I)
            if m_kab_doc:
                kab_candidate = m_kab_doc.group(1).strip()
                if kab_candidate.lower() not in ["ini", "tersebut", "terkait", "pamsimas"]:
                    kabupaten = kab_candidate
                    break

    # Jika masih belum ketemu, cari di nama file (baik pola "Kabupaten X" maupun nama kabupaten langsung)
    if not kabupaten:
        m_fn = re.search(r"Kabupaten\s+([A-Za-z]+)", clean_fn, re.I)
        if m_fn:
            kabupaten = m_fn.group(1).strip()
        else:
            # Cocokkan token nama file dengan master kabupaten
            fn_tokens = re.split(r"[-_\s\.]+", os.path.splitext(clean_fn)[0])
            for loc in KNOWN_LOCATIONS:
                k_clean = re.sub(r"^(?:Kabupaten|Kota)\s*", "", loc.get("kabupaten", ""), flags=re.I).strip()
                if any(t.lower() == k_clean.lower() for t in fn_tokens if len(t) > 2):
                    kabupaten = k_clean
                    break

    # Deteksi tahun dari teks atau nama file jika belum ada
    if not tahun or tahun == "2026":
        m_th = re.search(r"\b(202\d)\b", clean_fn) or re.search(r"\b(202\d)\b", " ".join(p.text for p in doc.paragraphs[:20]))
        if m_th:
            tahun = m_th.group(1)

    # 2. Nomor & Tanggal SPT
    nomor_spt = ""
    tanggal_spt_raw = ""
    for p in doc.paragraphs[:50]:
        t = p.text.strip()
        m_spt = re.search(r"\b(CK\d+[A-Za-z0-9\/\-\.]+202\d[A-Za-z0-9\/\-\.]*)\b", t)
        if m_spt and not nomor_spt:
            nomor_spt = m_spt.group(1).strip()
        m_tspt = re.search(
            r"tanggal\s+([0-9]{1,2}\s+[A-Za-z]+(?:\s+[0-9]{4})?\s*[-–—]\s*[0-9]{1,2}\s+[A-Za-z]+\s+[0-9]{4}|"
            r"[0-9]{1,2}\s*[-–—]\s*[0-9]{1,2}\s+[A-Za-z]+\s+[0-9]{4}|"
            r"[0-9]{1,2}\s+[A-Za-z]+\s+[0-9]{4})",
            t, re.I
        )
        if m_tspt and not tanggal_spt_raw:
            tanggal_spt_raw = m_tspt.group(1).strip()

    # 3. Tanggal Laporan & Dibuat di
    tempat_pembuatan = "Semarang"
    tanggal_laporan_raw = ""
    for p in doc.paragraphs[-35:]:
        t = p.text.strip()
        m_tempat = re.search(r"Dibuat di\s+([A-Za-z]+)", t, re.I)
        if m_tempat:
            tempat_pembuatan = m_tempat.group(1).strip()
        m_tlap = re.search(r"pada tanggal\s+([0-9]{1,2}\s+[A-Za-z]+\s+[0-9]{4})", t, re.I)
        if m_tlap:
            tanggal_laporan_raw = m_tlap.group(1).strip()

    tanggal_laporan_iso = parse_indo_date_to_iso(tanggal_laporan_raw)

    # 4. Tanggal Dinas (Mulai & Selesai)
    trip_start_iso = ""
    trip_end_iso = ""
    p_start, p_end = parse_date_range_str(tanggal_spt_raw)
    if not p_start:
        p_start, p_end = parse_date_range_str(clean_fn)
    if p_start:
        trip_start_iso = f"{p_start[0]:04d}-{p_start[1]:02d}-{p_start[2]:02d}"
    if p_end:
        trip_end_iso = f"{p_end[0]:04d}-{p_end[1]:02d}-{p_end[2]:02d}"

    tanggal_spt_iso = trip_start_iso or parse_indo_date_to_iso(tanggal_spt_raw)

    # Otomatisasi Tanggal Laporan: Jika belum ada di dokumen, jika dinas multi-hari diset H+1 hari akhir, atau H+1 Tanggal SPT
    auto_hplus1 = False
    if not tanggal_laporan_iso:
        if trip_end_iso and trip_start_iso and trip_end_iso != trip_start_iso:
            tanggal_laporan_iso = get_hplus1_iso(trip_end_iso)
            auto_hplus1 = True
        elif trip_end_iso or trip_start_iso:
            tanggal_laporan_iso = get_hplus1_iso(trip_end_iso or trip_start_iso)
            auto_hplus1 = True
        elif tanggal_spt_iso:
            tanggal_laporan_iso = get_hplus1_iso(tanggal_spt_iso)
            auto_hplus1 = True

    # 5. Tabel Pelaksana
    pelaksana = []
    has_pelaksana_table = False

    # Load master personel for clean matching
    master_personel = []
    personel_file = os.path.join(os.path.dirname(__file__), "..", "..", "data", "Nama Personel.txt")
    if os.path.exists(personel_file):
        with open(personel_file, "r", encoding="utf-8") as f:
            for line in f:
                cl = re.sub(r"^\d+\.\s*", "", line).strip()
                cl = normalize_personel_title(cl)
                if cl and cl not in master_personel:
                    master_personel.append(cl)

    for t in doc.tables:
        header_text = " ".join([c.text.lower() for c in t.rows[0].cells])
        if any(k in header_text for k in ["nama", "pelaksana", "tanda tangan", "ttd", "jabatan"]):
            has_pelaksana_table = True
            for r in t.rows[1:]:
                c_texts = [c.text.strip() for c in r.cells if c.text.strip()]
                for val in c_texts:
                    if not re.match(r"^\d+\.?$", val) and len(val) > 3 and not re.search(r"^\d+\.\s*$", val):
                        if val not in pelaksana and not any(k in val.lower() for k in ["tanda tangan", "jabatan", "nip"]):
                            # Cocokkan dengan master personel agar gelar/ejaan rapi
                            matched = None
                            for mp in master_personel:
                                if val.lower() in mp.lower() or mp.lower() in val.lower():
                                    matched = mp
                                    break
                            pelaksana.append(matched or normalize_personel_title(val))
                        break
        elif len(t.columns) == 4 and len(t.rows) >= 1:
            first_row = [c.text.strip() for c in t.rows[0].cells]
            if re.match(r"^\d+\.", first_row[0]) and len(first_row[1]) > 3:
                has_pelaksana_table = True
                for r in t.rows:
                    name_val = r.cells[1].text.strip()
                    if name_val and name_val not in pelaksana:
                        matched = None
                        for mp in master_personel:
                            if name_val.lower() in mp.lower() or mp.lower() in name_val.lower():
                                matched = mp
                                break
                        pelaksana.append(matched or normalize_personel_title(name_val))

    # 6. Parse Photos & Locations using docx_service
    parsed_photos = []
    detected_villages = []
    try:
        photos_raw, def_kab, def_prov, detected_villages, _ = parse_document(docx_path)
        for p in photos_raw:
            parsed_photos.append({
                "id": str(p.get("id", uuid.uuid4().hex[:8])),
                "filename": p.get("filename", ""),
                "url": p.get("url", ""),
                "width": p.get("width", 0),
                "height": p.get("height", 0),
                "desa": p.get("desa", ""),
                "kecamatan": p.get("kecamatan", ""),
                "dusun": p.get("dusun", ""),
                "timeStr": p.get("time_str", ""),
                "lat": p.get("lat", "-7.391520"),
                "lon": p.get("lon", "110.312050"),
                "kabupaten": p.get("kabupaten") or kabupaten,
                "provinsi": p.get("provinsi", "Jawa Tengah"),
                "has_timestamp": bool(p.get("has_timestamp", False)),
                "alreadyStamped": bool(p.get("has_timestamp", False)),
                "suggestedCrop": p.get("suggested_crop_bottom", 0.18)
            })
    except Exception as e:
        log_error("inspect_parse_document", e)

    # 7. Section B. Hasil yang Dicapai (Desa & Narasi)
    villages = []
    current_desa = None
    current_narrative_paras = []
    in_hasil_section = False

    for p in doc.paragraphs:
        txt = p.text.strip()
        if "B. HASIL YANG DICAPAI" in txt.upper():
            in_hasil_section = True
            continue
        if any(k in txt.upper() for k in ["C. KESIMPULAN", "PENUTUP", "DOKUMENTASI KEGIATAN"]):
            if current_desa:
                current_desa["narrative"] = "\n".join(current_narrative_paras).strip()
                current_desa["notes"] = current_desa["narrative"]
                villages.append(current_desa)
                current_desa = None
                current_narrative_paras = []
            break

        if in_hasil_section and txt:
            m_des = re.search(r"(?:^\d+[\.\)]\s*)?Desa\s+([A-Za-z0-9]+)(?:\s+(?:Kecamatan|Kec\.?)\s+([A-Za-z0-9]+))?", txt, re.I)
            is_heading = m_des and (len(txt.split()) <= 6 or p.style.name.startswith("Heading") or any(r.bold for r in p.runs[:2]))
            if is_heading:
                if current_desa:
                    current_desa["narrative"] = "\n".join(current_narrative_paras).strip()
                    current_desa["notes"] = current_desa["narrative"]
                    villages.append(current_desa)
                    current_narrative_paras = []
                desa_name = fuzzy_match_village(m_des.group(1).strip())
                kec_name = m_des.group(2).strip() if m_des.group(2) else ""
                if not kec_name:
                    _, _, looked_kec, _ = lookup_location(desa_name, "", kabupaten)
                    kec_name = looked_kec
                current_desa = {
                    "id": f"v_{desa_name.lower()}",
                    "desa": desa_name,
                    "kecamatan": kec_name,
                    "notes": "",
                    "narrative": ""
                }
            elif current_desa:
                current_narrative_paras.append(txt)

    if current_desa:
        current_desa["narrative"] = "\n".join(current_narrative_paras).strip()
        current_desa["notes"] = current_desa["narrative"]
        villages.append(current_desa)

    # Fallback desa jika belum ditemukan di Section B
    # 1. Dari token nama file yang cocok dengan master desa
    if not villages:
        fn_tokens = re.split(r"[-_\s\.]+", os.path.splitext(clean_fn)[0])
        for loc in KNOWN_LOCATIONS:
            d_name = loc.get("desa", "").strip()
            if d_name and any(t.lower() == d_name.lower() for t in fn_tokens if len(t) > 2):
                k_name = loc.get("kecamatan", "").strip()
                kb_name = re.sub(r"^(?:Kabupaten|Kota)\s*", "", loc.get("kabupaten", ""), flags=re.I).strip()
                villages.append({
                    "id": f"v_{d_name.lower()}",
                    "desa": d_name,
                    "kecamatan": k_name,
                    "notes": "",
                    "narrative": ""
                })
                if not kabupaten and kb_name:
                    kabupaten = kb_name
                break

    # 2. Dari detected_villages (hasil scan tabel stempel foto / parse_document)
    if not villages and detected_villages:
        for v_name in detected_villages:
            _, _, k_hint, kb_hint = lookup_location(v_name, "", kabupaten)
            villages.append({
                "id": f"v_{v_name.lower()}",
                "desa": v_name,
                "kecamatan": k_hint,
                "notes": "",
                "narrative": ""
            })
            if not kabupaten and kb_hint:
                kabupaten = re.sub(r"^(?:Kabupaten|Kota)\s*", "", kb_hint, flags=re.I).strip()

    # 3. Dari pencarian nama desa di seluruh paragraf dokumen
    if not villages:
        for p in doc.paragraphs:
            m_des_any = re.search(r"\bDesa\s+([A-Za-z0-9]+)", p.text, re.I)
            if m_des_any:
                cand_desa = fuzzy_match_village(m_des_any.group(1).strip())
                if cand_desa:
                    _, _, k_hint, kb_hint = lookup_location(cand_desa, "", kabupaten)
                    villages.append({
                        "id": f"v_{cand_desa.lower()}",
                        "desa": cand_desa,
                        "kecamatan": k_hint,
                        "notes": "",
                        "narrative": ""
                    })
                    if not kabupaten and kb_hint:
                        kabupaten = re.sub(r"^(?:Kabupaten|Kota)\s*", "", kb_hint, flags=re.I).strip()
                    break

    # Jika masih belum ada desa sama sekali, inisialisasi minimal 1 desa kosong
    if not villages:
        villages.append({
            "id": "v_init",
            "desa": "",
            "kecamatan": "",
            "notes": "",
            "narrative": ""
        })

    # 8. Checklist Audit Kelengkapan Dokumen
    audit_items = []

    # Item 1: Kabupaten
    if kabupaten:
        audit_items.append({"field": "kabupaten", "status": "complete", "label": "Kabupaten", "message": f"Terdeteksi: Kabupaten {kabupaten}"})
    else:
        audit_items.append({"field": "kabupaten", "status": "incomplete", "label": "Kabupaten", "message": "Belum terdeteksi di dokumen"})

    # Item 2: Nomor SPT
    if nomor_spt:
        audit_items.append({"field": "nomor_spt", "status": "complete", "label": "Nomor SPT", "message": f"Terdeteksi: {nomor_spt}"})
    else:
        audit_items.append({"field": "nomor_spt", "status": "incomplete", "label": "Nomor SPT", "message": "Belum tercantum di dokumen"})

    # Item 3: Tanggal SPT
    if tanggal_spt_iso:
        audit_items.append({"field": "tanggal_spt", "status": "complete", "label": "Tanggal SPT", "message": f"Terdeteksi: {tanggal_spt_raw}"})
    else:
        audit_items.append({"field": "tanggal_spt", "status": "incomplete", "label": "Tanggal SPT", "message": "Belum tercantum di dokumen"})

    # Item 4: Jadwal Perjalanan Dinas
    if trip_start_iso:
        trip_txt = f"{trip_start_iso} s/d {trip_end_iso}" if trip_end_iso and trip_end_iso != trip_start_iso else f"{trip_start_iso} (1 Hari)"
        audit_items.append({"field": "trip_dates", "status": "complete", "label": "Jadwal Dinas", "message": f"Terdeteksi: {trip_txt}"})
    else:
        audit_items.append({"field": "trip_dates", "status": "incomplete", "label": "Jadwal Dinas", "message": "Belum tercantum di dokumen"})

    # Item 5: Tanggal Laporan
    if auto_hplus1:
        audit_items.append({"field": "tanggal_laporan", "status": "incomplete", "label": "Tanggal Laporan", "message": f"Belum ada di dokumen (Otomatis diset default H+1 Dinas: {tanggal_laporan_iso})"})
    elif tanggal_laporan_iso:
        audit_items.append({"field": "tanggal_laporan", "status": "complete", "label": "Tanggal Laporan", "message": f"Terdeteksi: {tanggal_laporan_raw}"})
    else:
        audit_items.append({"field": "tanggal_laporan", "status": "incomplete", "label": "Tanggal Laporan", "message": "Belum tercantum di dokumen"})

    # Item 6: Tabel Personil Monev
    if has_pelaksana_table and len(pelaksana) > 0:
        audit_items.append({"field": "pelaksana", "status": "complete", "label": "Tabel Pelaksana", "message": f"Ditemukan tabel dengan {len(pelaksana)} personil pelaksana dinas"})
    else:
        audit_items.append({"field": "pelaksana", "status": "incomplete", "label": "Tabel Pelaksana", "message": "Belum ada tabel personil / tanda tangan pelaksana dinas"})

    # Item 7: Desa & Narasi Hasil yang Dicapai
    valid_villages = [v for v in villages if v.get("desa")]
    if len(valid_villages) > 0:
        has_narrative = any(len(v.get("narrative", "")) > 30 for v in valid_villages)
        v_names = ", ".join([v["desa"] for v in valid_villages[:3]]) + ("..." if len(valid_villages) > 3 else "")
        if has_narrative:
            audit_items.append({"field": "villages", "status": "complete", "label": "Desa & Narasi", "message": f"Ditemukan {len(valid_villages)} desa ({v_names}) beserta catatan capaian"})
        else:
            audit_items.append({"field": "villages", "status": "incomplete", "label": "Desa & Narasi", "message": f"Ditemukan {len(valid_villages)} desa ({v_names}), namun uraian narasi masih kosong"})
    else:
        audit_items.append({"field": "villages", "status": "incomplete", "label": "Desa & Narasi", "message": "Belum ada desa yang teridentifikasi"})

    # Item 8: Foto Dokumentasi Lapangan
    if len(parsed_photos) > 0:
        stamped_cnt = sum(1 for p in parsed_photos if p.get("has_timestamp"))
        if stamped_cnt == len(parsed_photos):
            audit_items.append({"field": "photos", "status": "complete", "label": "Foto Dokumentasi", "message": f"Ditemukan {len(parsed_photos)} foto (seluruh foto telah memiliki timestamp)"})
        elif stamped_cnt > 0:
            audit_items.append({"field": "photos", "status": "incomplete", "label": "Foto Dokumentasi", "message": f"Ditemukan {len(parsed_photos)} foto ({stamped_cnt} foto bertimestamp, {len(parsed_photos)-stamped_cnt} foto belum bertimestamp)"})
        else:
            audit_items.append({"field": "photos", "status": "incomplete", "label": "Foto Dokumentasi", "message": f"Ditemukan {len(parsed_photos)} foto di dokumen (BELUM memiliki timestamp GPS)"})
    else:
        audit_items.append({"field": "photos", "status": "incomplete", "label": "Foto Dokumentasi", "message": "Belum ada foto dokumentasi lapangan di dalam dokumen"})

    completed_cnt = sum(1 for a in audit_items if a["status"] == "complete")
    total_cnt = len(audit_items)
    score_pct = int((completed_cnt / total_cnt) * 100)
    summary_text = f"Dokumen berhasil dibuka ({score_pct}% Terisi). {completed_cnt} data terisi lengkap, {total_cnt - completed_cnt} data belum lengkap dan siap dilengkapi."

    return {
        "success": True,
        "filename": filename,
        "report_data": {
            "kabupaten": kabupaten,
            "tahun": tahun,
            "nomor_spt": nomor_spt,
            "tanggal_spt": tanggal_spt_iso,
            "tanggal_laporan": tanggal_laporan_iso,
            "tempat_pembuatan": tempat_pembuatan,
            "trip_start_date": trip_start_iso,
            "trip_end_date": trip_end_iso,
            "include_pelaksana": has_pelaksana_table and len(pelaksana) > 0,
            "pelaksana": pelaksana,
            "villages": villages
        },
        "photos": parsed_photos,
        "audit_report": {
            "score_percent": score_pct,
            "completed_count": completed_cnt,
            "incomplete_count": total_cnt - completed_cnt,
            "summary": summary_text,
            "items": audit_items
        }
    }
